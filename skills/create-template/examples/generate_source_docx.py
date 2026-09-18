#!/usr/bin/env python3
"""Create the Normal-style source DOCX used by the first-check example."""

import argparse
from pathlib import Path

from docx import Document


def build_fixture(destination: Path) -> None:
    """Write the minimal source document without direct body-run styling."""
    if destination.exists():
        raise FileExistsError(f"Refusing to overwrite existing file: {destination}")

    document = Document()
    heading = document.add_paragraph(style="Heading 1")
    heading.paragraph_format.page_break_before = True
    heading.add_run("Executive Summary")

    body = document.add_paragraph(style="Normal")
    body.add_run("Assessment results are summarized here.")

    document.core_properties.title = "Synthetic Create Template First Check"
    document.core_properties.author = "Ghostwriter skill fixture"
    document.save(destination)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("destination", type=Path)
    arguments = parser.parse_args()
    try:
        build_fixture(arguments.destination)
    except FileExistsError as error:
        parser.error(str(error))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
